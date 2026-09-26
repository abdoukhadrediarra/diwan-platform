import re
from docx import Document

def fix_tashkeel_lams(input_path, output_path):
    doc = Document(input_path)
    
    # PATTERN 1: Standard AL (ال)
    # Matches: [Optional prefixes] + Alif + Lam + [Moon Letter]
    pattern_al = re.compile(r'(?<![\u0621-\u064A])([وكفب]?[\u064B-\u065F]*)(ا)(ل)([أإآابجحخعغفقكمهوي])')
    
    # PATTERN 2: LI-L (لِل) preposition + definite article
    # Matches: [Optional wa/fa] + First Lam & diacritics + Second Lam + [Moon Letter]
    # The negative lookahead (?!ه) prevents it from mistakenly adding a sukun to unvocalized "لله"
    pattern_lil = re.compile(r'(?<![\u0621-\u064A])([وف]?[\u064B-\u065F]*ل[\u064B-\u065F]*)(ل)(?!ه[\u064B-\u065F]*(?:\s|[.,؛،]|$))([أإآابجحخعغفقكمهوي])')

    def process_text(text):
        # Apply Pattern 1 (Standard Alif-Lam)
        # \1\2\3ْ\4 inserts the captured groups and adds Sukun (\u0652) to the Lam
        text = pattern_al.sub(r'\1\2\3ْ\4', text)
        
        # Apply Pattern 2 (Li-L construct)
        # \1\2ْ\3 inserts Group 1 (لِ), Group 2 (ل) + Sukun (\u0652), Group 3 (Moon Letter)
        text = pattern_lil.sub(r'\1\2ْ\3', text)
        
        return text

    # Process standard paragraphs
    for para in doc.paragraphs:
        for run in para.runs:
            if run.text:
                # We check and process the text inside each run
                run.text = process_text(run.text)
                
    # Process text inside tables (if your diwans use grid layouts)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    for run in para.runs:
                        if run.text:
                            run.text = process_text(run.text)

    # Save the corrected document
    doc.save(output_path)
    print(f"Correction complete. Saved as {output_path}")

# Run the function
if __name__ == "__main__":
    # Replace with your actual file names
    input_file = "tools\diwan_1.docx"
    output_file = "diwan_1_fixed.docx"
    
    fix_tashkeel_lams(input_file, output_file)