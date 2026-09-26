"""
apps/corpus/parsing.py

Turns ONE poem into a plain Python dict (the JSON that goes into PostgreSQL).
It never touches the database. Three ways to use it:

    web page      python tools/poem_json_web/app.py            (drop .docx files, get JSON)
    command line  python apps/corpus/parsing.py corpus-source/diwan-01/ --out json/
    Django        python manage.py import_poems json/          (or the .docx files directly)

How a poem is read (empty paragraphs are ignored):
    first text(s) before the abyat                     -> section "muqaddima", kind "prose"
    the second text, right before the abyat            -> section "title",     kind "title"
      = the poem's own NAME given by the author (only when there are at least two texts)
    a single text before the abyat                     -> section "title" too, but only when the
      abyat spell it: then that line is the poem's acrostic name, not an opening text
    BOLD text anywhere in the opening, before the first bayt -> not split off at all: the name is
      read straight out of the muqaddima paragraph it sits in, which is kept exactly as written.
      This is for poems whose manuscript never repeats the name as its own line — duplicating it
      onto a new line would add words the author never wrote there.
    a short line naming a letter ("اللام"), among the abyat -> section "matn", kind "header"
      = a GROUPED acrostic: every bayt that follows, until the next header, starts with that
        letter. A shadda in the name (الحقّ) means the letter opens two groups in a row.
        Section headers are not counted among the "texts before the abyat" above.
    paragraphs containing "|"                          -> section "matn",      kind "bayt"
    paragraphs after the last bayt                     -> section "khatima",   kind "prose"

The poem's name (the "title" field):
    title_source "bold"        text marked bold anywhere before the first bayt (checked first)
    title_source "name_line"   the second text before the abyat, as above; or the single text
                               before the abyat when the abyat spell it
    title_source "first_sadr"  everything else: the poem is known by the sadr of its first bayt

A name found in bold is always treated as the poem's acrostic name (is_acrostic is always true),
since the reviewer marked it deliberately — a low letter-match only earns a warning, in case the
wrong span was highlighted, never a downgrade to "not acrostic".

The name is what the platform works with (display, search, months, events, days).
Whether an own name is ALSO an acrostic is a secondary fact, checked leniently: its letters are
compared with the first letters of the opening abyat, and if at least half of them are found
there, in order, is_acrostic is true. A poem that spells its name several times, or goes on
with other acrostics afterwards, still counts. A name that is not an acrostic
(e.g. هَذِهِ أَسْمَاءُ السُّوَرِ) is normal and gives no warning.
"""
import argparse
import difflib
import hashlib
import json
import re
import sys
from collections import Counter
from pathlib import Path

import docx  # pip install python-docx

try:
    from .arabic import acrostic_letters, acrostic_letters_by_group, bayt_initial, clean_display_text, header_letter
    from .transcription import transcribe_hemistichs
except ImportError:  # run as a standalone script
    sys.path.insert(0, str(Path(__file__).resolve().parent))
    from arabic import acrostic_letters, acrostic_letters_by_group, bayt_initial, clean_display_text, header_letter
    from transcription import transcribe_hemistichs

CODE_RE = re.compile(r"^D(\d{2})K(\d{2,3})", re.IGNORECASE)
ORPHAN_MARK_RE = re.compile(r"(?:^|\s)[\u064B-\u0652\u0670]+")  # haraka with no letter under it
DOUBLE_MARK_RE = re.compile(r"[\u064B-\u0650\u0652]\u0651?[\u064B-\u0650\u0652]")  # two vowels on one letter
ACROSTIC_MAJORITY = 0.5


def code_from_name(name: str) -> tuple[str, int, int]:
    match = CODE_RE.match(Path(name).stem.strip())
    if not match:
        raise ValueError(f"{name}: the name must start with a code like D01K08")
    diwan_number, poem_number = int(match.group(1)), int(match.group(2))
    return f"D{diwan_number:02d}K{poem_number:02d}", diwan_number, poem_number


def acrostic_match(phrase: str, sadrs: list[str]) -> int:
    """Percentage of the phrase's letters found, in order, at the start of the opening abyat."""
    expected = acrostic_letters(phrase)
    if not expected or not sadrs:
        return 0
    initials = "".join(bayt_initial(s) for s in sadrs)
    window = initials[: len(expected) + max(3, len(expected) // 5)]  # room for a few extra or missing abyat
    matcher = difflib.SequenceMatcher(None, expected, window, autojunk=False)
    found = sum(block.size for block in matcher.get_matching_blocks())
    return round(100 * found / len(expected))


def _is_red(run) -> bool:
    color = run.font.color
    if color is None or color.type is None or color.rgb is None:
        return False
    r, g, b = color.rgb
    return r >= 0xC0 and g < 0x50 and b < 0x50


LETTER_RE = re.compile(r"[\u0621-\u064A]")


def red_marks(paragraph) -> list[dict]:
    """Red text in a paragraph, grouped into pieces, with the word(s) around each piece.

    A piece with at least one red LETTER is a real flag. A piece that is only harakat
    (e.g. a kasra left red after the word was recoloured black) is reported as a leftover
    and does not count. Red spaces are invisible and ignored.
    """
    text, spans = "", []
    for run in paragraph.runs:
        if run.text and _is_red(run):
            if spans and spans[-1][1] == len(text):
                spans[-1][1] += len(run.text)
            else:
                spans.append([len(text), len(text) + len(run.text)])
        text += run.text

    marks = []
    for start, end in spans:
        piece = text[start:end]
        if not piece.strip():
            continue
        while start > 0 and not text[start - 1].isspace() and text[start - 1] != "|":
            start -= 1
        while end < len(text) and not text[end].isspace() and text[end] != "|":
            end += 1
        marks.append({"red": piece.strip(), "word": text[start:end].strip(), "is_flag": bool(LETTER_RE.search(piece))})
    return marks


def bold_spans(paragraph) -> list[str]:
    """The paragraph's explicitly bold-formatted text, as separate spans in reading order (more
    than one when plain text sits between two bolded parts). A style that happens to render bold
    does not count — only a deliberate bold toggle on the run does, since that is the reviewer's
    own mark."""
    spans, current = [], ""
    for run in paragraph.runs:
        if run.text and run.bold:
            current += run.text
        elif current:
            spans.append(current)
            current = ""
    if current:
        spans.append(current)
    return [span.strip() for span in spans if span.strip()]


def parse_poem_docx(source, filename: str | None = None) -> dict:
    """source: a path, or an uploaded file object (then give its filename)."""
    name = filename or Path(source).name
    code_from_name(name)  # fail early on a bad file name
    document = docx.Document(source)
    paragraphs = [(p.text, red_marks(p), bold_spans(p)) for p in document.paragraphs]
    notes = ["the file contains tables; only normal paragraphs are read"] if document.tables else []
    return parse_paragraphs(name, paragraphs, notes)


def parse_poem_text(code: str, text: str) -> dict:
    """Pasted text: one paragraph per line (plain text carries no bold formatting)."""
    return parse_paragraphs(code, [(line, [], []) for line in text.splitlines()], [])


def parse_paragraphs(name: str, raw_paragraphs, warnings: list[str]) -> dict:
    code, diwan_number, poem_number = code_from_name(name)
    warnings = list(warnings)

    paragraphs = []  # (text, red marks, bold text)
    for text, red, bold in raw_paragraphs:
        text = re.sub(r"\s+", " ", clean_display_text(text)).strip()
        if text:
            paragraphs.append((text, red, bold))

    bayt_indexes = [i for i, (text, _, _) in enumerate(paragraphs) if "|" in text]
    if not bayt_indexes:
        raise ValueError(f"{code}: no bayt found (no line contains '|' between the two hemistichs)")
    first_bayt, last_bayt = bayt_indexes[0], bayt_indexes[-1]
    opening_sadrs = [paragraphs[k][0].split("|")[0].strip() for k in bayt_indexes]

    # Section headers ("اللام"...) are recognized wherever they appear among the abyat, and are
    # not counted among the "texts before the abyat" used to find the poem's name below.
    headers = {i: header_letter(text) for i, (text, _, _) in enumerate(paragraphs)
              if i <= last_bayt and header_letter(text)}
    opening_indexes = [i for i in range(first_bayt) if i not in headers]

    # A reviewer sometimes marks the poem's own name in BOLD, right inside the opening prose,
    # instead of writing it again as a separate line — many manuscripts never repeat it, so the
    # digitisation must not invent that repetition either. This is checked first: found or not,
    # it settles the name outright, and the opening text it sits in is kept exactly as written,
    # unsplit, in the muqaddima.
    bold_phrase = " ".join(span for i in range(first_bayt) for span in paragraphs[i][2]).strip()

    # The poem's name, when there is one, is the second (non-header) text, right before the abyat.
    # When there is only one, it is the name as well if the abyat spell it (an acrostic name) —
    # unless section headers are already present, which confirms it without that fragile check.
    if bold_phrase:
        title_index = None      # the name stays inside its muqaddima paragraph; nothing is split off
    elif len(opening_indexes) >= 2:
        title_index = opening_indexes[-1]
    elif len(opening_indexes) == 1:
        title_index = opening_indexes[0]
        if not headers:
            if acrostic_match(paragraphs[title_index][0], opening_sadrs) < 100 * ACROSTIC_MAJORITY:
                title_index = None
    else:
        title_index = None
    if not bold_phrase and len(opening_indexes) > 2:
        warnings.append(f"{len(opening_indexes)} texts before the abyat (section headers aside): "
                        f"line {opening_indexes[-1] + 1} was taken as the poem's name, please confirm")

    lines, bayt_number, has_open_flags = [], 0, False
    for i, (text, red, bold) in enumerate(paragraphs):
        position = i + 1
        if i == title_index:
            section, kind = "title", "title"
        elif i in headers:
            section, kind = "matn", "header"
        elif i < first_bayt:
            section, kind = "muqaddima", "prose"
        elif i > last_bayt:
            section, kind = "khatima", "prose"
        elif "|" in text:
            section, kind = "matn", "bayt"
        else:
            section, kind = "matn", "prose"
            warnings.append(f"line {position}: text without '|' between two abyat (kept as prose inside the matn)")

        # The exact text(s) that were bolded in THIS paragraph, cleaned the same way as the line
        # itself, so the front-end can find and colour them red inside the untouched prose — kept
        # as separate spans (there can be more than one, with plain words between them), and only
        # for the paragraph(s) that actually make up the bold-detected name, never a stray bold
        # word elsewhere (a khatima emphasis, say) that has nothing to do with the acrostic.
        acrostic_spans = []
        if bold_phrase and section == "muqaddima" and bold:
            acrostic_spans = [s for s in (re.sub(r"\s+", " ", clean_display_text(b)).strip() for b in bold) if s]

        if kind == "bayt":
            bayt_number += 1
            label = f"bayt {bayt_number}"
            hemistichs = [h.strip() for h in text.split("|")]
            if len(hemistichs) not in (2, 4) or not all(hemistichs):
                warnings.append(f"{label}: expected 2 or 4 non-empty hemistichs, found {len(hemistichs)}")
        else:
            label = f"{section} (line {position})"
            hemistichs = [text]

        for m in ORPHAN_MARK_RE.finditer(text):
            context = text[max(0, m.start() - 15): m.end() + 8].strip()
            warnings.append(f"{label}: haraka with no letter under it near «{context}»")
        for m in DOUBLE_MARK_RE.finditer(text):
            start, end = text.rfind(" ", 0, m.start()) + 1, text.find(" ", m.end())
            warnings.append(f"{label}: two vowels on one letter in «{text[start:end if end >= 0 else len(text)]}»")
        for mark in red:
            if mark["is_flag"]:
                has_open_flags = True
                warnings.append(f"{label}: red-flagged «{mark['word']}»")
            else:
                warnings.append(f"{label}: a red haraka is left on «{mark['word']}» (its letters are black); "
                                f"not counted as a flag, but recolour it before printing")

        lines.append({
            "position": position,
            "section": section,
            "kind": kind,
            "bayt_number": bayt_number if kind == "bayt" else None,
            "hemistichs": hemistichs,
            "transcription": {"local": transcribe_hemistichs(hemistichs)},   # generated; not part of content_hash
            "acrostic_spans": acrostic_spans,  # the bolded substring(s) of this line, if any; not part of content_hash
        })

    abyat = [l for l in lines if l["kind"] == "bayt"]
    sadrs = [l["hemistichs"][0] for l in abyat]

    if bold_phrase:
        # the reviewer's own mark settles it: this text is the name, acrostic or not, exactly as
        # bolded, and it is used as-is even if the letters do not line up well with the abyat —
        # a low match is flagged, in case the wrong span was highlighted, but never overridden
        title, title_source = bold_phrase, "bold"
        match = acrostic_match(title, sadrs)
        is_acrostic = True
        if match < 100 * ACROSTIC_MAJORITY:
            warnings.append(f"name: the bold text «{title}» only matches {match}% of the abyat's "
                            f"opening letters; please check that the right span was highlighted")
    elif headers:
        # a grouped acrostic: check every bayt against the header of its own group, and compare
        # the sequence of header letters with the name's letters (a shadda counts twice).
        current_letter = None
        for line in lines:
            if line["kind"] == "header":
                current_letter = headers[line["position"] - 1]
            elif line["kind"] == "bayt" and current_letter:
                initial = bayt_initial(line["hemistichs"][0])
                if initial != current_letter:
                    warnings.append(f"bayt {line['bayt_number']}: starts with «{initial}», "
                                    f"expected «{current_letter}» for this group")
        header_sequence = "".join(headers[i] for i in sorted(headers))
        if title_index is not None:
            title, title_source = paragraphs[title_index][0], "name_line"
            expected_sequence = acrostic_letters_by_group(title)
            matcher = difflib.SequenceMatcher(None, expected_sequence, header_sequence, autojunk=False)
            found = sum(block.size for block in matcher.get_matching_blocks())
            match = round(100 * found / len(expected_sequence)) if expected_sequence else 0
            is_acrostic = match >= 100 * ACROSTIC_MAJORITY
            if not is_acrostic:
                warnings.append(f"acrostic: the section headers spell «{header_sequence}», only "
                                f"{match}% of the name's letters «{expected_sequence}» found there")
        else:                               # section headers with no name line above them: unusual
            title, title_source, is_acrostic, match = sadrs[0], "first_sadr", False, None
            warnings.append("acrostic: section headers were found, but no name line introduces them; "
                            "please check that this poem's own name has not been left out")
    elif title_index is not None:
        title, title_source = paragraphs[title_index][0], "name_line"
        match = acrostic_match(title, sadrs)
        is_acrostic = match >= 100 * ACROSTIC_MAJORITY
    else:                                   # no own name: the poem is known by its first sadr
        title, title_source, is_acrostic, match = sadrs[0], "first_sadr", False, None

    sizes = Counter(len(l["hemistichs"]) for l in abyat)
    if len(sizes) > 1:
        warnings.append(f"abyat have different numbers of hemistichs: {dict(sizes)}")

    hash_source = json.dumps([[l["section"], l["kind"], l["hemistichs"]] for l in lines], ensure_ascii=False)
    return {
        "code": code,
        "diwan": diwan_number,
        "number": poem_number,
        "source_file": name,
        "incipit": sadrs[0],
        "hemistichs_per_bayt": sizes.most_common(1)[0][0],
        "bayt_count": len(abyat),
        "title": title,
        "title_source": title_source,
        "is_acrostic": is_acrostic,
        "acrostic_match": match,
        "counts": {s: sum(1 for l in lines if l["section"] == s) for s in ("muqaddima", "title", "matn", "khatima")},
        "has_open_flags": has_open_flags,
        "content_hash": hashlib.sha256(hash_source.encode("utf-8")).hexdigest(),
        "warnings": warnings,
        "lines": lines,
    }


REQUIRED_KEYS = {"code", "diwan", "number", "incipit", "hemistichs_per_bayt", "bayt_count", "title",
                 "title_source", "is_acrostic", "has_open_flags", "content_hash", "lines"}


def load_poem_json(path) -> dict:
    """Read a JSON file produced by this module (e.g. downloaded from the web page)."""
    return check_poem_json(json.loads(Path(path).read_text(encoding="utf-8")), Path(path).name)


def check_poem_json(data: dict, name: str) -> dict:
    """Check that a decoded JSON object is a poem made by this module."""
    if not isinstance(data, dict):
        raise ValueError(f"{name}: not a poem JSON")
    missing = REQUIRED_KEYS - data.keys()
    if missing:
        raise ValueError(f"{name}: not a poem JSON (missing {', '.join(sorted(missing))}); make it again with the web page")
    data.setdefault("warnings", [])
    data.setdefault("acrostic_match", None)
    data.setdefault("source_file", name)
    return data


def name_summary(data: dict) -> str:
    if data["title_source"] == "first_sadr":
        return f"named by its first sadr «{data['title']}»"
    acrostic = f"acrostic {data['acrostic_match']}%" if data["is_acrostic"] else "not acrostic"
    return f"name «{data['title']}» ({acrostic})"


def find_poem_files(paths, suffixes=(".docx", ".json")):
    files = []
    for p in map(Path, paths):
        if p.is_dir():
            # look inside sub-folders too, so corpus-json/diwan-01/… is found from corpus-json/
            files += sorted(f for f in p.rglob("*") if f.suffix.lower() in suffixes and not f.name.startswith("~$"))
        elif p.suffix.lower() in suffixes:
            files.append(p)
        else:
            raise ValueError(f"not a {' / '.join(suffixes)} file or a folder: {p}")
    return files


if __name__ == "__main__":
    ap = argparse.ArgumentParser(description="Parse poem .docx files into JSON (no database).")
    ap.add_argument("paths", nargs="+")
    ap.add_argument("--out", help="folder where one <code>.json per poem is written")
    args = ap.parse_args()
    for f in find_poem_files(args.paths, suffixes=(".docx",)):
        data = parse_poem_docx(f)
        text = json.dumps(data, ensure_ascii=False, indent=2)
        print(f"{data['code']}: {name_summary(data)}, {data['counts']['muqaddima']} muqaddima, "
              f"{data['bayt_count']} abyat, {data['counts']['khatima']} khatima, {len(data['warnings'])} warning(s)",
              file=sys.stderr)
        for w in data["warnings"]:
            print(f"  ! {w}", file=sys.stderr)
        if args.out:
            Path(args.out).mkdir(parents=True, exist_ok=True)
            (Path(args.out) / f"{data['code']}.json").write_text(text, encoding="utf-8")
        else:
            print(text)
